# -*- coding: utf-8 -*-
"""
Created on Tue Apr  7 13:04:27 2026

@author: aengstrom
"""
from openpyxl import load_workbook
import calendar
import os
import logging
from pathlib import Path
import pandas as pd
from autogc_validation.database.conn import connection
from autogc_validation.database.operations import get_canister_periods
from autogc_validation.database.enums import COLUMN_CALIBRANTS, ColumnType, SampleTypeLetter
logger = logging.getLogger(__name__)
def populate_mdvr(mdvr_path: os.PathLike,
                  month: int,
                  year: int,
                  db_path: os.PathLike = None,
                  site_id: int = None,
                  calibrants: dict[ColumnType, int] = COLUMN_CALIBRANTS) -> None:
    add_month_to_mdvr(mdvr_path, month, year)
    if db_path and site_id:
        last_day = calendar.monthrange(year, month)[1]
        start_date = f"{year}-{month:02d}-01 00:00"
        end_date = f"{year}-{month:02d}-{last_day:02d} 23:59"
        add_site_to_mdvr(mdvr_path, db_path, site_id)
        add_active_qc(mdvr_path, db_path, site_id, start_date, end_date, calibrants)


DATE_ROW = 5
DATE_COL = 2
def add_month_to_mdvr(mdvr_path: os.PathLike, month: int, year: int) -> None:
    mdvr_path = Path(mdvr_path)
    month_name = calendar.month_name[month]
    wb = load_workbook(mdvr_path)
    ws = wb["Validator Notes"]
    
    ws.cell(DATE_ROW, DATE_COL).value = f"{month_name} {year}"
    wb.save(mdvr_path)
    logger.info(f"Wrote date to MDVR doc: '{month_name} {year}'")

SITE_ROW = 4
SITE_COL = 2    
def add_site_to_mdvr(mdvr_path: os.PathLike, database, site_id) -> None:
    mdvr_path = Path(mdvr_path)
    sql = """SELECT name_short, name_long
    FROM sites
    WHERE site_id = ?"""
    with connection(database) as conn:
        cursor = conn.execute(sql, (site_id,))
        row = cursor.fetchone()
    if row is None:
        logger.warning("No site found for site_id=%s", site_id)
        return
    short, long = row
    wb = load_workbook(mdvr_path)
    ws = wb["Validator Notes"]
    ws.cell(SITE_ROW, SITE_COL).value = f"{short} - {long}"
    wb.save(mdvr_path)
    logger.info(f"Wrote site to MDVR doc: '{short} - {long}'")
    
STD_ROW = 6
STD_COL = 1
PRIMARY_CAN_COL = 2
DIL_RATIO_COL = 3
START_COL = 4
PLOT_TARGET_COL = 13
BP_TARGET_COL = 14
def add_active_qc(mdvr_path: os.PathLike, 
                  database,
                  site_id: int,
                  start_date: str,
                  end_date: str,
                  calibrants: dict[ColumnType, int] = COLUMN_CALIBRANTS
                  ) -> None:
    mdvr_path = Path(mdvr_path)
    standards = ["CVS", "LCS", "RTS"]
    row_idx = STD_ROW
    wb = load_workbook(mdvr_path)
    ws = wb["Reference Info"]
    for standard in standards:
        active = get_canister_periods(database = database,
                                      site_id = site_id,
                                      canister_type = standard,
                                      start_date = start_date,
                                      end_date = end_date,
                                      output_unit = "ppbc",
                                      include_install_info = True)
        
        
        plot_calibrant = calibrants[ColumnType("PLOT")]
        bp_calibrant = calibrants[ColumnType("BP")]
        available = set(active.columns)
        if plot_calibrant not in available:
            logger.warning("PLOT calibrant AQS %s not found in %s canister — cell will be blank", plot_calibrant, standard)
            plot_calibrant = None
        if bp_calibrant not in available:
            logger.warning("BP calibrant AQS %s not found in %s canister — cell will be blank", bp_calibrant, standard)
            bp_calibrant = None

        for index, row in active.iterrows():
            start = index
            primary_can = row["primary_canister_id"]
            dil_ratio = row["dilution_ratio"]
            plot_target = row[plot_calibrant] if plot_calibrant is not None else None
            bp_target = row[bp_calibrant] if bp_calibrant is not None else None

            
            ws.cell(row_idx, STD_COL).value = standard
            ws.cell(row_idx, PRIMARY_CAN_COL).value = primary_can
            ws.cell(row_idx, DIL_RATIO_COL).value = dil_ratio
            ws.cell(row_idx, START_COL).value = start
            ws.cell(row_idx, PLOT_TARGET_COL).value = plot_target
            ws.cell(row_idx, BP_TARGET_COL).value = bp_target

            row_idx += 1
    wb.save(mdvr_path)
    logger.info("Wrote QC standards to MDVR doc for site_id=%s", site_id)
#CVS data start cells
CVS_DATE_ROW, CVS_DATE_COL = 4, 1
PLOT_CVS_ROW, PLOT_CVS_COL = 4, 2
BP_CVS_ROW, BP_CVS_COL = 4, 5
#LCS data start cells
LCS_DATE_ROW, LCS_DATE_COL = 4, 8
PLOT_LCS_ROW, PLOT_LCS_COL = 4, 9
BP_LCS_ROW, BP_LCS_COL = 4, 10
#RTS data start cells
RTS_DATE_ROW, RTS_DATE_COL = 19, 8
PLOT_RTS_ROW, PLOT_RTS_COL = 19, 9
BP_RTS_ROW, BP_RTS_COL = 19, 10

# Maps SampleTypeLetter to (start_row, date_col, plot_col, bp_col)
_QC_COORDS = {
    SampleTypeLetter.CVS: (CVS_DATE_ROW, CVS_DATE_COL, PLOT_CVS_COL, BP_CVS_COL),
    SampleTypeLetter.LCS: (LCS_DATE_ROW, LCS_DATE_COL, PLOT_LCS_COL, BP_LCS_COL),
    SampleTypeLetter.RTS: (RTS_DATE_ROW, RTS_DATE_COL, PLOT_RTS_COL, BP_RTS_COL),
}

def add_qc_to_mdvr(mdvr_path: os.PathLike,
                   qc_df: pd.DataFrame,
                   calibrants: dict[ColumnType, int] = COLUMN_CALIBRANTS) -> None:
    mdvr_path = Path(mdvr_path)
    wb = load_workbook(mdvr_path)
    ws = wb["Calculations"]
    qc_type = qc_df.attrs["sample_type"]
    if qc_type not in _QC_COORDS:
        raise ValueError(
            f"add_qc_to_mdvr: unsupported sample type {qc_type!r}. "
            f"Expected one of {[k.name for k in _QC_COORDS]}"
        )
    start_row, date_col, plot_col, bp_col = _QC_COORDS[qc_type]
    plot_aqs = calibrants[ColumnType("PLOT")]
    bp_aqs = calibrants[ColumnType("BP")]
    available = set(qc_df.columns)
    if plot_aqs not in available:
        logger.warning("PLOT calibrant AQS %s not found in %s DataFrame — cell will be blank", plot_aqs, qc_type.name)
        plot_aqs = None
    if bp_aqs not in available:
        logger.warning("BP calibrant AQS %s not found in %s DataFrame — cell will be blank", bp_aqs, qc_type.name)
        bp_aqs = None

    for i, (index, row) in enumerate(qc_df.iterrows()):
        ws.cell(start_row + i, date_col).value = index
        ws.cell(start_row + i, plot_col).value = row[plot_aqs] if plot_aqs is not None else None
        ws.cell(start_row + i, bp_col).value = row[bp_aqs] if bp_aqs is not None else None

    wb.save(mdvr_path)
    logger.info("Wrote %s QC data to MDVR Calculations sheet (%d rows)", qc_type.name, len(qc_df))
    
    
    
    
    
    